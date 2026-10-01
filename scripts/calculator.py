#!/usr/bin/env python3
"""
calculator.py — a deal calculator workbook for a negotiation.

The negotiation cheat-sheet prompt asks the model to end its output with one
fenced block describing how the role sheet scores a deal:

    ```calculator
    {"title": "Moms.com (seller)",
     "unit": "$",
     "inputs": [
       {"name": "fee",  "label": "Licensing fee per episode", "type": "number", "default": 50000},
       {"name": "runs", "label": "Runs per episode", "type": "choice", "default": "5",
        "options": [{"label": "4", "value": 500000}, {"label": "5", "value": 250000},
                    {"label": "6", "value": 0}]}
     ],
     "outputs": [
       {"name": "revenue", "label": "Program revenue", "formula": "fee*100"},
       {"name": "net",     "label": "Net revenue",     "formula": "revenue+runs"}
     ],
     "reference": {"score": "net", "reservation": 2500000, "target": 4500000},
     "packages": [{"label": "Anchor", "values": {"fee": 75000, "runs": "5"}}]}
    ```

`extract()` takes that block out of the Markdown (so it never reaches the Word
file) and `build()` turns it into "Calculator - <Case>.xlsx": three package
columns, yellow input cells, dropdowns for choice issues, live formulas, and
the distance from your reservation value and target in red or green.

A formula may use input and earlier output names, numbers, + - * / ^ ( ) and
comparison operators, and IF, MIN, MAX, ROUND, ABS, SUM, AND, OR. Anything else
is refused: a wrong calculator in the room is worse than none.
"""

import json
import re
from pathlib import Path

FENCE_RE = re.compile(r"\n*(?:^|\n)#{0,3}\s*(?:Calculator spec[^\n]*\n+)?```calculator\s*\n(.*?)\n```[ \t]*\n?",
                      re.DOTALL | re.IGNORECASE)
NAME_RE = re.compile(r"[A-Za-z_][A-Za-z0-9_]*")
FUNCS = {"IF", "MIN", "MAX", "ROUND", "ABS", "SUM", "AND", "OR"}
ALLOWED_CHARS_RE = re.compile(r"^[A-Za-z0-9_+\-*/^().,<>=%\s]*$")
MAX_PACKAGES = 3
PREFIX = "Calculator - "


class CalculatorError(ValueError):
    pass


# ── Markdown ──────────────────────────────────────────────────────────────────

def extract(md: str) -> "tuple[str, dict | None]":
    """(markdown without the calculator block, spec or None)."""
    m = FENCE_RE.search(md or "")
    if not m:
        return md, None
    cleaned = (md[:m.start()].rstrip() + "\n\n" + md[m.end():].lstrip()).strip() + "\n"
    try:
        spec = json.loads(m.group(1))
    except ValueError:
        return cleaned, None
    return cleaned, spec if isinstance(spec, dict) else None


# ── Validation and evaluation ─────────────────────────────────────────────────

def _names(spec: dict) -> "tuple[list, list]":
    inputs = spec.get("inputs") or []
    outputs = spec.get("outputs") or []
    if not inputs or not outputs:
        raise CalculatorError("needs at least one input and one output")
    seen = set()
    for item in inputs + outputs:
        n = item.get("name", "")
        if not NAME_RE.fullmatch(n) or n.upper() in FUNCS or n in seen:
            raise CalculatorError(f"bad or repeated name: {n!r}")
        seen.add(n)
    for i in inputs:
        if i.get("type") == "choice":
            opts = i.get("options") or []
            if not opts or any(not isinstance(o.get("value"), (int, float)) for o in opts):
                raise CalculatorError(f"choice {i['name']!r} needs options with numeric values")
        elif not isinstance(i.get("default", 0), (int, float)):
            raise CalculatorError(f"number {i['name']!r} needs a numeric default")
    return inputs, outputs


def validate(spec: dict) -> None:
    """Raise CalculatorError unless every formula is safe and evaluates."""
    inputs, outputs = _names(spec)
    known = {i["name"] for i in inputs}
    for o in outputs:
        f = str(o.get("formula", "")).lstrip("=")
        if not f or not ALLOWED_CHARS_RE.match(f):
            raise CalculatorError(f"formula for {o['name']!r} has characters outside the whitelist")
        for tok in NAME_RE.findall(f):
            if tok.upper() in FUNCS or tok in known:
                continue
            raise CalculatorError(f"formula for {o['name']!r} uses unknown name {tok!r}")
        known.add(o["name"])
    ref = spec.get("reference") or {}
    if ref.get("score") and ref["score"] not in {o["name"] for o in outputs}:
        raise CalculatorError("reference.score must name an output")
    evaluate(spec, {})                      # must compute with the defaults


def _default_value(i: dict):
    if i.get("type") == "choice":
        labels = [str(o["label"]) for o in i["options"]]
        d = str(i.get("default", labels[0]))
        return d if d in labels else labels[0]
    return i.get("default", 0)


def _py(formula: str) -> str:
    f = formula.lstrip("=").replace("^", "**").replace("<>", "!=")
    f = re.sub(r"(?<![<>!=])=(?!=)", "==", f)
    return re.sub(r"\b(IF|MIN|MAX|ROUND|ABS|SUM|AND|OR)\b", lambda m: "_" + m.group(1).upper(), f,
                  flags=re.IGNORECASE)


_ENV = {
    "_IF": lambda c, a, b=0: a if c else b, "_MIN": min, "_MAX": max,
    "_ROUND": lambda x, n=0: round(x, int(n)), "_ABS": abs,
    "_SUM": lambda *a: sum(a), "_AND": lambda *a: all(a), "_OR": lambda *a: any(a),
}


def evaluate(spec: dict, values: dict) -> dict:
    """Outputs for one package (missing inputs take their defaults). Mirrors Excel."""
    inputs, outputs = _names(spec)
    scope = {}
    for i in inputs:
        v = values.get(i["name"], _default_value(i))
        if i.get("type") == "choice":
            table = {str(o["label"]): o["value"] for o in i["options"]}
            if str(v) not in table:
                raise CalculatorError(f"{i['name']}: {v!r} is not one of {list(table)}")
            v = table[str(v)]
        scope[i["name"]] = v
    out = {}
    for o in outputs:
        try:
            scope[o["name"]] = out[o["name"]] = eval(_py(str(o["formula"])), {"__builtins__": {}}, {**_ENV, **scope})
        except Exception as e:
            raise CalculatorError(f"formula for {o['name']!r} does not evaluate: {e}")
    return out


# ── Workbook ──────────────────────────────────────────────────────────────────

def filename(case_title: str) -> str:
    safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", case_title or "Negotiation").strip() or "Negotiation"
    return f"{PREFIX}{safe[:80]}.xlsx"


def build(spec: dict, out_path: Path) -> Path:
    """Write the workbook. Raises CalculatorError on a spec that fails validate()."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    validate(spec)
    inputs, outputs = _names(spec)
    packages = (spec.get("packages") or [])[:MAX_PACKAGES]
    while len(packages) < MAX_PACKAGES:
        packages.append({"label": f"Package {chr(65 + len(packages))}", "values": {}})
    unit = spec.get("unit") or ""
    money = f'"{unit}"#,##0;[Red]-"{unit}"#,##0' if unit in ("$", "€", "£") else "#,##0.00"

    wb = Workbook()
    ws = wb.active
    ws.title = "Deal"
    lk = wb.create_sheet("Lookup")
    lk["A1"], lk["B1"] = "Choice lists (do not edit)", None

    navy, yellow, grey = "1F3864", "FFF2CC", "F2F2F2"
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    cols = [get_column_letter(2 + k) for k in range(MAX_PACKAGES)]          # B, C, D

    ws["A1"] = spec.get("title") or "Deal calculator"
    ws["A1"].font = Font(bold=True, size=14, color=navy)
    ws["A2"] = "Change the yellow cells. Everything else recalculates."
    ws["A2"].font = Font(italic=True, color="595959")
    hdr = 4
    ws.cell(hdr, 1, "Inputs").font = Font(bold=True, color="FFFFFF")
    ws.cell(hdr, 1).fill = PatternFill("solid", fgColor=navy)
    for k, c in enumerate(cols):
        cell = ws[f"{c}{hdr}"]
        cell.value = str(packages[k].get("label") or f"Package {chr(65 + k)}")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.cell(hdr, 5, "Notes").font = Font(bold=True, color="FFFFFF")
    ws.cell(hdr, 5).fill = PatternFill("solid", fgColor=navy)

    # where each name lives, per package column: name -> {col: expression}
    ref: dict = {}
    row, lk_col = hdr + 1, 1
    for i in inputs:
        ws.cell(row, 1, i.get("label") or i["name"])
        ws.cell(row, 5, i.get("note") or "")
        ref[i["name"]] = {}
        if i.get("type") == "choice":
            opts = i["options"]
            a, b = get_column_letter(lk_col), get_column_letter(lk_col + 1)
            lk[f"{a}2"], lk[f"{b}2"] = i["name"], "value"
            for n, o in enumerate(opts):
                lk[f"{a}{3 + n}"], lk[f"{b}{3 + n}"] = str(o["label"]), o["value"]
            lo, hi = 3, 2 + len(opts)
            dv = DataValidation(type="list", formula1=f"=Lookup!${a}${lo}:${a}${hi}", allow_blank=False)
            ws.add_data_validation(dv)
            lk_col += 3
        for k, c in enumerate(cols):
            cell = ws[f"{c}{row}"]
            v = (packages[k].get("values") or {}).get(i["name"], _default_value(i))
            cell.fill, cell.border = PatternFill("solid", fgColor=yellow), box
            cell.alignment = Alignment(horizontal="right")
            if i.get("type") == "choice":
                labels = [str(o["label"]) for o in i["options"]]
                cell.value = str(v) if str(v) in labels else labels[0]
                dv.add(cell)
                ref[i["name"]][c] = (f"INDEX(Lookup!${b}${lo}:${b}${hi},"
                                     f"MATCH({c}{row}&\"\",Lookup!${a}${lo}:${a}${hi},0))")
            else:
                cell.value = v if isinstance(v, (int, float)) else i.get("default", 0)
                cell.number_format = i.get("format") or ("0%" if i.get("percent") else "#,##0.##")
                ref[i["name"]][c] = f"{c}{row}"
        row += 1

    row += 1
    ws.cell(row, 1, "Results").font = Font(bold=True, color="FFFFFF")
    for c in ["A"] + cols + ["E"]:
        ws[f"{c}{row}"].fill = PatternFill("solid", fgColor=navy)
    row += 1

    def xl(formula: str, c: str) -> str:
        f = formula.lstrip("=")
        def sub(m):
            tok = m.group(0)
            if tok.upper() in FUNCS:
                return tok.upper()
            return f"({ref[tok][c]})" if tok in ref else tok
        return "=" + NAME_RE.sub(sub, f)

    score_name = (spec.get("reference") or {}).get("score") or outputs[-1]["name"]
    for o in outputs:
        ws.cell(row, 1, o.get("label") or o["name"])
        ws.cell(row, 5, o.get("note") or "")
        ref[o["name"]] = {}
        for c in cols:
            cell = ws[f"{c}{row}"]
            cell.value = xl(str(o["formula"]), c)
            cell.number_format = o.get("format") or money
            cell.border, cell.fill = box, PatternFill("solid", fgColor=grey)
            ref[o["name"]][c] = f"{c}{row}"
        if o["name"] == score_name:
            for c in ["A"] + cols:
                ws[f"{c}{row}"].font = Font(bold=True)
        row += 1

    reference = spec.get("reference") or {}
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    for key, label in (("reservation", "reservation value"), ("target", "target")):
        if not isinstance(reference.get(key), (int, float)):
            continue
        row += 1
        ws.cell(row, 1, f"My {label}")
        for c in cols:
            ws[f"{c}{row}"].value = reference[key]
            ws[f"{c}{row}"].number_format = money
        row += 1
        ws.cell(row, 1, f"Deal vs {label}").font = Font(bold=True)
        rng = f"{cols[0]}{row}:{cols[-1]}{row}"
        for c in cols:
            cell = ws[f"{c}{row}"]
            cell.value = f"={ref[score_name][c]}-{c}{row - 1}"
            cell.number_format, cell.font, cell.border = money, Font(bold=True), box
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=red))
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=green))
    if reference.get("batna_note"):
        row += 2
        ws.cell(row, 1, f"BATNA: {reference['batna_note']}").font = Font(italic=True)

    ws.column_dimensions["A"].width = 44
    for c in cols:
        ws.column_dimensions[c].width = 20
    ws.column_dimensions["E"].width = 60
    ws.freeze_panes = f"B{hdr + 1}"
    lk.sheet_state = "hidden"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
