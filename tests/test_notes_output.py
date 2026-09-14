import sys

import canvas_refresh as cr
import cheat_sheet


def test_markdown_to_docx_smoke(tmp_path):
    from docx import Document

    out = tmp_path / "notes.docx"
    cr.markdown_to_docx(
        "# H1\n\n- bullet **bold**\n\n1. one\n\n---\n\nplain *it*\n",
        out, "Title", {"Generated": "now"},
    )
    doc = Document(str(out))
    texts  = [p.text for p in doc.paragraphs]
    styles = [p.style.name for p in doc.paragraphs]

    assert "Title" in texts
    assert "Generated: now" in texts
    assert "H1" in texts and styles[texts.index("H1")] == "Heading 1"
    bullet = doc.paragraphs[texts.index("bullet bold")]
    assert bullet.style.name == "List Bullet"
    assert any(r.bold for r in bullet.runs)
    assert styles[texts.index("one")] == "List Number"
    italic = doc.paragraphs[texts.index("plain it")]
    assert any(r.italic for r in italic.runs)


def test_cheat_sheet_delegates_to_generate_notes(patch_courses, monkeypatch, capsys):
    posting = {"id": 1, "name": "LTV | Class 3", "description": "<p>q</p>",
               "submission_types": ["not_graded"], "due_at": "2026-09-02T12:00:00Z"}
    monkeypatch.setattr(cr, "assignments_on", lambda cid, date_str, kinds=cr.SESSION_KINDS: [posting])

    seen = {}
    monkeypatch.setattr(cr, "generate_notes", lambda session: seen.update(session))
    monkeypatch.setattr(sys, "argv", ["cheat_sheet.py", "260902", "LTV"])

    cheat_sheet.main()

    assert seen["abbrev"] == "LTV" and seen["date_str"] == "260902"
    assert seen["course_id"] == 1
    assert seen["assignments"] == [posting] and seen["assignment"] is posting
    assert "found: LTV | Class 3" in capsys.readouterr().out


def test_cheat_sheet_rejects_unknown_course(patch_courses, monkeypatch):
    monkeypatch.setattr(sys, "argv", ["cheat_sheet.py", "260902", "NOPE"])
    try:
        cheat_sheet.main()
    except SystemExit as e:
        assert "Unknown course 'NOPE'" in str(e.code)
    else:
        raise AssertionError("expected SystemExit")
