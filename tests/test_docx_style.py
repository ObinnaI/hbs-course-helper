from docx import Document

import canvas_refresh as cr


MD = """# Cheat Sheet: Pave (A)

## The case in 90 seconds

**Pave** is a **$1.6B** compensation-data startup. Plain *italic* and `code`.

### Q1. What should Matt do?

**The 20-second answer, say this first:** Raise now.

| Metric | Value | Page |
|---|---|---|
| ARR | **$12M** | p. 4 |
| Burn | $1.1M/mo | p. 6 |

- bullet one
  - nested
1) numbered
> a quoted line
"""


def test_cheatsheet_style_renders_tables_and_title_block(tmp_path, monkeypatch):
    monkeypatch.setenv("NOTES_STYLE", "cheatsheet")
    out = tmp_path / "x.docx"
    cr.markdown_to_docx(MD, out, "Launching Tech Ventures: September 16, 2026",
                        {"Generated": "now", "Course": "LTV"})
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]

    assert texts[0] == "Case Discussion Preparation"
    assert texts[1] == "Launching Tech Ventures: September 16, 2026"
    assert "Generated: now" in texts
    # The model's own "# Cheat Sheet:" H1 is not repeated under the title block.
    assert not any(t.startswith("Cheat Sheet:") for t in texts)

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3 and len(table.columns) == 3
    assert table.cell(0, 0).text == "Metric"
    assert table.cell(1, 1).text == "$12M"
    assert any(r.bold for r in table.cell(0, 1).paragraphs[0].runs)

    body = next(p for p in doc.paragraphs if p.text.startswith("Pave is a"))
    bold = [r.text for r in body.runs if r.bold]
    assert bold == ["Pave", "$1.6B"]
    assert all(r.font.name == "Times New Roman" for r in body.runs if r.font.name)
    assert all(r.font.size.pt == 10 for r in body.runs)

    h1 = next(p for p in doc.paragraphs if p.text == "The case in 90 seconds")
    assert h1.style.name == "Heading 2"
    assert all(r.bold for r in h1.runs)

    styles = [p.style.name for p in doc.paragraphs]
    assert "List Bullet" in styles and "List Bullet 2" in styles and "List Number" in styles
    quote = next(p for p in doc.paragraphs if p.text == "a quoted line")
    assert all(r.italic for r in quote.runs)


def test_compact_style_keeps_old_look(tmp_path, monkeypatch):
    monkeypatch.setenv("NOTES_STYLE", "compact")
    out = tmp_path / "y.docx"
    cr.markdown_to_docx(MD, out, "Title", {"Generated": "now"})
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Title"
    assert "Cheat Sheet: Pave (A)" in texts        # compact keeps the H1
    body = next(p for p in doc.paragraphs if p.text.startswith("Pave is a"))
    assert all(r.font.size.pt == 12 for r in body.runs)
