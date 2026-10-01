import json

import pytest

import calculator as calc
import canvas_common as cc
import canvas_refresh as cr

MOMS = {
    "title": "Moms.com (seller)", "unit": "$",
    "inputs": [
        {"name": "fee", "label": "Licensing fee per episode", "type": "number", "default": 50000},
        {"name": "runs", "label": "Runs per episode", "type": "choice", "default": "5",
         "options": [{"label": "4", "value": 500000}, {"label": "5", "value": 250000},
                     {"label": "6", "value": 0}, {"label": "7", "value": -250000},
                     {"label": "8", "value": -500000}]},
        {"name": "y1", "label": "Share paid in year 1", "type": "number", "default": 0.3},
        {"name": "y2", "label": "Share paid in year 2", "type": "number", "default": 0.2},
        {"name": "y3", "label": "Share paid in year 3", "type": "number", "default": 0},
    ],
    "outputs": [
        {"name": "revenue", "label": "Program revenue", "formula": "fee*100"},
        {"name": "financing", "label": "Financing cost", "formula": "revenue*(y1*0.20+y2*0.35+y3*0.50)"},
        {"name": "net", "label": "Net revenue", "formula": "revenue-financing+runs"},
        {"name": "ok", "label": "Clears the floor?", "formula": "IF(fee>=35000, 1, 0)"},
    ],
    "reference": {"score": "net", "reservation": 2500000, "target": 4500000},
    "packages": [{"label": "Role-sheet example", "values": {}},
                 {"label": "Anchor", "values": {"fee": 75000, "runs": "5", "y1": 0.5, "y2": 0, "y3": 0}}],
}


def test_extract_strips_the_block():
    md = "# Cheat Sheet: X\n\nbody\n\n## Calculator spec\n```calculator\n" + json.dumps(MOMS) + "\n```\n"
    out, spec = calc.extract(md)
    assert "calculator" not in out.lower() and out.strip().endswith("body")
    assert spec["title"] == "Moms.com (seller)"
    assert calc.extract("# no block\n") == ("# no block\n", None)
    out, spec = calc.extract("x\n```calculator\n{not json}\n```\n")
    assert spec is None and "```" not in out


def test_evaluate_reproduces_the_role_sheet_example():
    out = calc.evaluate(MOMS, {})
    assert out["revenue"] == 5_000_000
    assert round(out["financing"]) == 650_000
    assert round(out["net"]) == 4_600_000            # → $2.1M over the $2.5M alternative
    assert out["ok"] == 1
    assert round(calc.evaluate(MOMS, {"fee": 45000, "runs": "8", "y1": 0.3, "y2": 0.3, "y3": 0.4})["net"]) == 2_357_500


@pytest.mark.parametrize("formula", ["__import__('os').system('x')", "fee*unknown", "fee;1", "A1+fee", "fee*'x'"])
def test_unsafe_or_unknown_formulas_are_refused(formula):
    bad = json.loads(json.dumps(MOMS))
    bad["outputs"][0]["formula"] = formula
    with pytest.raises(calc.CalculatorError):
        calc.validate(bad)


def test_workbook_has_inputs_formulas_and_dropdowns(tmp_path):
    from openpyxl import load_workbook
    out = calc.build(MOMS, tmp_path / calc.filename("Moms.com I"))
    assert out.name == "Calculator - Moms.com I.xlsx"
    wb = load_workbook(out)
    ws = wb["Deal"]
    cells = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
    assert ws.cell(cells["Licensing fee per episode"], 2).value == 50000
    assert ws.cell(cells["Licensing fee per episode"], 3).value == 75000        # second package
    assert ws.cell(cells["Runs per episode"], 2).value == "5"
    rev = ws.cell(cells["Program revenue"], 2).value
    assert rev == f"=(B{cells['Licensing fee per episode']})*100"
    net = ws.cell(cells["Net revenue"], 2).value
    assert "INDEX(Lookup!" in net and f"B{cells['Program revenue']}" in net
    assert ws.cell(cells["Clears the floor?"], 2).value.startswith("=IF(")
    assert ws.cell(cells["Deal vs reservation value"], 2).value == f"=B{cells['Net revenue']}-B{cells['My reservation value']}"
    assert ws.data_validations.dataValidation and wb["Lookup"].sheet_state == "hidden"
    assert ws.conditional_formatting


def test_calculator_is_not_a_reading(tmp_path):
    d = tmp_path / "261001 Viking"; d.mkdir()
    (d / "Viking_SandyWoodEDIT.pdf").write_bytes(b"%PDF")
    calc.build(MOMS, d / calc.filename("Viking"))
    assert cc.is_generated_file("Calculator - Viking.xlsx") and not cc.is_generated_file("Model.xlsx")
    assert [f.name for f in cr._reading_files(d)] == ["Viking_SandyWoodEDIT.pdf"]
