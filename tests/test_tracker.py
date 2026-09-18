import openpyxl

import participation_tracker as pt


def _courses(*abbrevs, term="Fall"):
    return {a: {"canvas_id": i + 1, "full_name": f"Full {a}", "folder_name": a,
                "folder_path": object(), "refinement_prompt": None, "term": term,
                "term_end": None}
            for i, a in enumerate(abbrevs)}


def test_course_order_is_sorted_and_needs_a_folder(monkeypatch):
    courses = _courses("LTV", "CFO", "AAA")
    courses["ZZZ"] = {"canvas_id": 9, "full_name": "No folder", "folder_path": None, "term": "Fall"}
    courses["SPR"] = {"canvas_id": 8, "full_name": "Spring one", "folder_path": object(), "term": "Spring"}
    monkeypatch.setattr(pt, "_COURSES", courses)
    assert pt.course_order("Fall") == ["AAA", "CFO", "LTV"]
    assert pt.course_order("Spring") == ["SPR"]
    assert pt.course_order() == ["AAA", "CFO", "LTV", "SPR"]


def test_colours_cycle_through_palette():
    order = [f"C{i}" for i in range(len(pt._PALETTE) + 2)]
    colours = pt.course_colors(order)
    assert colours["C0"] == pt._PALETTE[0]
    assert colours[order[len(pt._PALETTE)]] == pt._PALETTE[0]
    assert colours[order[1]] != colours[order[2]]


def _sessions_for(order):
    return {a: [{"name": f"{a} | Class 1: Case A", "due_at": "2026-09-02T12:00:00Z"},
                {"name": f"{a} | Class 2: Case B", "due_at": "2026-09-09T12:00:00Z"}]
            for a in order}


def test_ratings_survive_a_course_inserted_before_them(monkeypatch, tmp_path):
    out = tmp_path / "Fall" / "Participation Tracker.xlsx"
    monkeypatch.setattr(pt, "tracker_path", lambda term: out)
    monkeypatch.setattr(pt, "LEGACY_FILE", tmp_path / "none.xlsx")

    # Build with two courses, rate an LTV session.
    monkeypatch.setattr(pt, "_COURSES", _courses("CFO", "LTV"))
    monkeypatch.setattr(pt.path_config, "COURSE_NAMES", {"CFO": "Full CFO", "LTV": "Full LTV"})
    monkeypatch.setattr(pt, "COURSE_NAMES", {"CFO": "Full CFO", "LTV": "Full LTV"})
    monkeypatch.setattr(pt, "get_all_sessions", lambda order: _sessions_for(order))
    pt.build_tracker("Fall")

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    ltv_col = pt.course_start_col(1)          # CFO is 0, LTV is 1
    assert ws.cell(row=1, column=ltv_col).value == "Full LTV"
    ws.cell(row=pt.DATA_START_ROW, column=ltv_col + 2, value="great")
    wb.save(out)

    # Add AAA, which sorts first and shifts every column group right.
    monkeypatch.setattr(pt, "_COURSES", _courses("CFO", "LTV", "AAA"))
    names = {"AAA": "Full AAA", "CFO": "Full CFO", "LTV": "Full LTV"}
    monkeypatch.setattr(pt, "COURSE_NAMES", names)
    pt.build_tracker("Fall")

    ws = openpyxl.load_workbook(out).active
    ltv_col = pt.course_start_col(2)
    assert ws.cell(row=1, column=ltv_col).value == "Full LTV"
    assert ws.cell(row=pt.DATA_START_ROW, column=ltv_col + 2).value == "great"
    aaa_col = pt.course_start_col(0)
    assert ws.cell(row=pt.DATA_START_ROW, column=aaa_col + 2).value in (None, "")

    ratings = pt.read_existing_ratings(out, ["AAA", "CFO", "LTV"])
    assert ratings == {"AAA": {}, "CFO": {}, "LTV": {"260902": "great"}}


def test_extract_case_title():
    assert pt.extract_case_title("CFO | Class 3: The DCF Method") == "The DCF Method"
    assert pt.extract_case_title("CATS | Class 1 | Capitalism") == "Capitalism"
    assert pt.extract_case_title("Plain title") == "Plain title"


def test_refresh_builds_one_sheet_per_live_term(monkeypatch, tmp_path):
    from datetime import datetime, timezone
    monkeypatch.setattr(pt, "DEST_ROOT", tmp_path)
    monkeypatch.setattr(pt, "LEGACY_FILE", tmp_path / "Participation Tracker.xlsx")
    courses = _courses("MP", "NEG", term="Fall")
    courses.update({a: dict(d, canvas_id=d["canvas_id"] + 10) for a, d in _courses("OLD", term="Summer").items()})
    monkeypatch.setattr(pt, "_COURSES", courses)
    monkeypatch.setattr(pt, "COURSE_NAMES", {a: f"Full {a}" for a in courses})

    def sessions(order):
        if "OLD" in order:   # a term that ended months ago
            return {"OLD": [{"name": "Class 1: X", "due_at": "2026-05-01T12:00:00Z"}]}
        return {a: [{"name": f"Class 1: {a}", "due_at": "2026-09-30T12:00:00Z"}] for a in order}
    monkeypatch.setattr(pt, "get_all_sessions", sessions)

    pt.refresh(now=datetime(2026, 9, 18, tzinfo=timezone.utc))
    assert (tmp_path / "Fall" / "Participation Tracker.xlsx").exists()
    assert not (tmp_path / "Summer" / "Participation Tracker.xlsx").exists()


def test_first_term_sheet_inherits_legacy_root_ratings(monkeypatch, tmp_path):
    monkeypatch.setattr(pt, "DEST_ROOT", tmp_path)
    legacy = tmp_path / "Participation Tracker.xlsx"
    monkeypatch.setattr(pt, "LEGACY_FILE", legacy)
    monkeypatch.setattr(pt, "tracker_path", lambda term: tmp_path / term / "Participation Tracker.xlsx")
    monkeypatch.setattr(pt, "_COURSES", _courses("LTV"))
    monkeypatch.setattr(pt, "COURSE_NAMES", {"LTV": "Full LTV"})
    monkeypatch.setattr(pt, "get_all_sessions", lambda order: _sessions_for(order))

    # Build a legacy-style sheet at the root and rate a session in it.
    monkeypatch.setattr(pt, "tracker_path", lambda term: legacy)
    pt.build_tracker("Fall")
    wb = openpyxl.load_workbook(legacy); ws = wb.active
    ws.cell(row=pt.DATA_START_ROW, column=3, value="good"); wb.save(legacy)

    monkeypatch.setattr(pt, "tracker_path", lambda term: tmp_path / term / "Participation Tracker.xlsx")
    pt.build_tracker("Fall")
    out = tmp_path / "Fall" / "Participation Tracker.xlsx"
    assert pt.read_existing_ratings(out, ["LTV"]) == {"LTV": {"260902": "good"}}
